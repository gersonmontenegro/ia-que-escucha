"""Genera el primer borrador de diapositivas y su versión de texto."""

from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = Path(__file__).resolve().parent
ARCHITECTURE = ".archify/architecture-rag-conferencia-20261013-223433/arquitectura-rag.html"

SLIDES = [
    {
        "title": "¿Cómo enseñarle a una IA a escuchar?",
        "minutes": "Apertura",
        "layout": "cover",
        "subtitle": "Exposición: 35 min · Preguntas: 10 min\nRAG sobre video: transcripción, búsqueda semántica y respuestas con minuto citado.",
    },
    {
        "title": "0–3 · Una pregunta y un minuto exacto",
        "minutes": "0–3 · Gancho",
        "layout": "content",
        "columns": [[
            "Demo en vivo: preguntar sobre el video de Saturno.",
            "Ver la respuesta y los fragmentos que la sustentan.",
            "Elegir una cita y saltar al momento del video.",
            "Pregunta: ¿qué acaban de descubrir en Saturno?",
        ]],
    },
    {
        "title": "3–7 · ¿Qué cambia cuando añadimos RAG?",
        "minutes": "3–7 · El problema",
        "layout": "two",
        "columns": [
            [
                "SIN RAG",
                "El modelo responde desde sus parámetros y conocimientos previos.",
                "Una noticia reciente puede ser desconocida o dar pie a una invención.",
                "No hay vínculo verificable con el video.",
            ],
            [
                "CON RAG",
                "Buscamos fragmentos de la transcripción antes de generar.",
                "El modelo recibe evidencia concreta y cita los fragmentos [n].",
                "Si el video no lo menciona, la respuesta debe decirlo.",
            ],
        ],
    },
    {
        "title": "7–10 · Arquitectura: dos recorridos",
        "minutes": "7–10 · Arquitectura",
        "layout": "two",
        "columns": [
            [
                "INGESTA · preparar el video",
                "MP4 → Whisper → segmentos con marcas de tiempo.",
                "Segmentos → chunks → embeddings → Qdrant.",
                "La transcripción queda en caché para reutilizarla.",
            ],
            [
                "CONSULTA · responder la pregunta",
                "Pregunta → embedding → búsqueda top-k en Qdrant.",
                "Qwen recibe pregunta y fragmentos recuperados.",
                "Respuesta con citas → reproductor en el minuto citado.",
                f"Mapa interactivo Archify: {ARCHITECTURE}",
            ],
        ],
    },
    {
        "title": "10–14 · Whisper convierte audio en evidencia temporal",
        "minutes": "10–14 · Whisper",
        "layout": "content",
        "columns": [[
            "Entrada: video MP4. Salida: segmentos {start, end, text}.",
            "Las marcas de tiempo permiten regresar al segundo relevante.",
            "faster-whisper en Docker (CPU) o MLX nativo en Apple Silicon.",
            "En video.mp4 (10:16): 123,9 s faster-whisper; 17,9 s MLX con modelo en caché.",
            "La primera ejecución de MLX tardó 81,0 s e incluyó descarga del modelo.",
        ]],
    },
    {
        "title": "14–18 · Chunking: contexto sin perder el tiempo",
        "minutes": "14–18 · Chunking",
        "layout": "two",
        "columns": [
            [
                "SI SON MUY CORTOS",
                "Cada vector contiene poco contexto.",
                "Una idea puede quedar partida entre resultados.",
            ],
            [
                "SI SON MUY LARGOS",
                "Se mezclan temas y baja la precisión del fragmento recuperado.",
                "El prompt incluye texto irrelevante.",
            ],
        ],
        "tail": [
            "Configuración de la demo: máximo 700 caracteres y solapamiento de 1 segmento.",
            "Cada chunk conserva el inicio y final para citar el minuto.",
        ],
    },
    {
        "title": "18–22 · Embeddings: significado como vector",
        "minutes": "18–22 · Embeddings",
        "layout": "content",
        "columns": [[
            "bge-m3 convierte cada fragmento y cada pregunta en un vector de 1024 dimensiones.",
            "Textos semánticamente parecidos quedan próximos en el espacio vectorial.",
            "La pregunta y los chunks deben usar el mismo modelo de embeddings.",
            "Un embedding representa texto; no es una respuesta ni una base de datos.",
        ]],
    },
    {
        "title": "22–26 · Búsqueda: NumPy explica, Qdrant escala",
        "minutes": "22–26 · Búsqueda",
        "layout": "two",
        "columns": [
            [
                "NUMPY · a mano",
                "Normalizamos vectores y calculamos producto punto.",
                "Ordenamos por similitud coseno y tomamos los mejores k.",
                "Sirve para ver qué operación hace la búsqueda.",
            ],
            [
                "QDRANT · servicio vectorial",
                "Guarda vectores y metadatos: texto, video y marcas de tiempo.",
                "Recupera los 4 fragmentos más similares en esta demo.",
                "La puntuación ordena resultados; no es una probabilidad de verdad.",
            ],
        ],
    },
    {
        "title": "26–28 · Generación: qué recibe el modelo",
        "minutes": "26–28 · Prompt",
        "layout": "content",
        "columns": [[
            "Instrucción: responde solo con la información de los fragmentos.",
            "Cita los fragmentos [n] y reconoce cuando falta evidencia.",
            "Prompt = instrucción + fragmentos recuperados + pregunta; modelo: Qwen3.8 27B.",
        ]],
    },
    {
        "title": "28–30 · Una respuesta que vuelve al video",
        "minutes": "28–30 · Citas",
        "layout": "content",
        "columns": [[
            "Los fragmentos conservan texto, video, puntuación y tiempo.",
            "La cita [n] lleva al momento correspondiente del reproductor.",
            "Medimos tiempo total, tokens de salida y tokens/s promedio (incluye espera inicial).",
        ]],
    },
    {
        "title": "30–32 · ¿Cuándo falla el RAG?",
        "minutes": "30–32 · Límites",
        "layout": "two",
        "columns": [
            [
                "RECUPERACIÓN",
                "Nombres propios o cifras pueden no recuperarse bien.",
                "Un chunk mal cortado puede perder contexto.",
                "Similitud alta no garantiza una respuesta.",
            ],
            [
                "GENERACIÓN",
                "La pregunta puede estar fuera del video.",
                "El modelo aún puede inferir de más.",
                "Inspeccionamos los fragmentos antes de confiar.",
            ],
        ],
    },
    {
        "title": "32–33 · Mejoras para explorar",
        "minutes": "32–33 · Mejoras",
        "layout": "content",
        "columns": [[
            "Híbrida: significado + palabras exactas.",
            "Reranking: reordena candidatos; HyDE ayuda a formular la búsqueda.",
            "Son mejoras para probar, no evidencia nueva.",
        ]],
    },
    {
        "title": "33–35 · Local, nube y frameworks",
        "minutes": "33–35 · Ecosistema",
        "layout": "two",
        "columns": [
            [
                "ESTA DEMO",
                "Ollama local para embeddings y chat.",
                "Qdrant y Gradio en Docker.",
                "Proveedores intercambiables por base URL, modelo y clave en .env.",
            ],
            [
                "OTRAS OPCIONES",
                "OpenAI-compatible permite cambiar a un proveedor en la nube.",
                "LangChain y LlamaIndex ofrecen componentes y abstracciones para RAG.",
                "Aquí escribimos el pipeline para hacer visibles sus etapas.",
            ],
        ],
    },
    {
        "title": "35–45 · Cierre y preguntas",
        "minutes": "35–45 · Preguntas",
        "layout": "content",
        "columns": [[
            "Cierre (35–36): una respuesta útil recupera evidencia y conserva su origen.",
            "Repo: transcripción → chunks → embeddings → búsqueda → generación.",
            "Después: retos en docs/retos.md.",
            "Preguntas del público (36–45).",
        ]],
    },
]


def add_bullets(placeholder, lines: list[str]) -> None:
    frame = placeholder.text_frame
    frame.clear()
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.level = 0


def build_presentation() -> None:
    presentation = Presentation()
    presentation.core_properties.title = "¿Cómo enseñarle a una IA a escuchar?"
    presentation.core_properties.subject = "Primer borrador: conferencia sobre RAG de video"

    for item in SLIDES:
        if item["layout"] == "cover":
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = item["title"]
            slide.placeholders[1].text = item["subtitle"]
            continue

        layout_index = 3 if item["layout"] == "two" else 1
        slide = presentation.slides.add_slide(presentation.slide_layouts[layout_index])
        slide.shapes.title.text = item["title"]

        if item["layout"] == "two":
            add_bullets(slide.placeholders[1], item["columns"][0])
            add_bullets(slide.placeholders[2], item["columns"][1] + item.get("tail", []))
        else:
            add_bullets(slide.placeholders[1], item["columns"][0])

    presentation.save(OUTPUT / "conferencia-rag-video.pptx")


def write_markdown() -> None:
    lines = [
        "# ¿Cómo enseñarle a una IA a escuchar?",
        "",
        "Iteración enfocada en 35 minutos de exposición y 10 minutos de preguntas. El PPTX usa layouts y placeholders estándar, sin tema gráfico fijado.",
        "",
    ]
    for number, item in enumerate(SLIDES, start=1):
        lines.extend([f"## {number}. {item['title']} ({item['minutes']})", ""])
        if item["layout"] == "cover":
            lines.extend([item["subtitle"], ""])
        elif item["layout"] == "two":
            columns = [item["columns"][0], item["columns"][1] + item.get("tail", [])]
            for heading, column in zip(("Columna izquierda", "Columna derecha"), columns):
                lines.append(f"**{heading}**")
                lines.extend(f"- {line}" for line in column)
                lines.append("")
        else:
            lines.extend(f"- {line}" for line in item["columns"][0])
            lines.append("")
    (OUTPUT / "contenido.md").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    build_presentation()
    write_markdown()
